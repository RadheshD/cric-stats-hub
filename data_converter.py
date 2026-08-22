import pandas as pd
import json
import os
from file_manager import read_json, save_json
from openpyxl.styles import Font, Alignment, Border, Side

def display_tables_from_json(file_path):
    """Convert JSON data into clean pandas DataFrames and print them."""
    data = read_json(file_path)
    if not data:
        return

    print("\n--- Match Summary ---")
    summary_df = pd.DataFrame([data['match_details']])
    print(summary_df.to_string(index=False))

    print("\n--- Team Scores ---")
    team_scores_df = pd.DataFrame(data['team_scores']).T
    print(team_scores_df.to_string())

    print("\n--- Batting Stats (India) ---")
    ind_bat_df = pd.DataFrame(data['batting_stats']['India'])
    print(ind_bat_df.to_string(index=False))

    print("\n--- Batting Stats (Australia) ---")
    aus_bat_df = pd.DataFrame(data['batting_stats']['Australia'])
    print(aus_bat_df.to_string(index=False))

    print("\n--- Bowling Stats (Australia) ---")
    aus_bowl_df = pd.DataFrame(data['bowling_stats']['Australia'])
    print(aus_bowl_df.to_string(index=False))

    print("\n--- Bowling Stats (India) ---")
    ind_bowl_df = pd.DataFrame(data['bowling_stats']['India'])
    print(ind_bowl_df.to_string(index=False))

def json_to_excel(json_file, excel_file):
    """Convert JSON data to Excel with multiple sheets and formatting."""
    data = read_json(json_file)
    if not data:
        return

    # Prepare dataframes
    summary_df = pd.DataFrame([data['match_details']])
    team_scores_df = pd.DataFrame(data['team_scores']).T.reset_index().rename(columns={'index': 'Team'})
    ind_bat_df = pd.DataFrame(data['batting_stats']['India'])
    aus_bat_df = pd.DataFrame(data['batting_stats']['Australia'])
    aus_bowl_df = pd.DataFrame(data['bowling_stats']['Australia'])
    ind_bowl_df = pd.DataFrame(data['bowling_stats']['India'])

    # Write to Excel
    with pd.ExcelWriter(excel_file, engine='openpyxl') as writer:
        summary_df.to_excel(writer, sheet_name='Match Summary', index=False)
        team_scores_df.to_excel(writer, sheet_name='Team Scores', index=False)
        
        # Combine batting stats
        ind_bat_df['Team'] = 'India'
        aus_bat_df['Team'] = 'Australia'
        batting_df = pd.concat([ind_bat_df, aus_bat_df], ignore_index=True)
        batting_df.to_excel(writer, sheet_name='Batting Scorecard', index=False)
        
        # Combine bowling stats
        aus_bowl_df['Team'] = 'Australia'
        ind_bowl_df['Team'] = 'India'
        bowling_df = pd.concat([aus_bowl_df, ind_bowl_df], ignore_index=True)
        bowling_df.to_excel(writer, sheet_name='Bowling Scorecard', index=False)

    # Format Excel file
    from openpyxl import load_workbook
    wb = load_workbook(excel_file)
    header_font = Font(bold=True)
    border_style = Side(style='thin')
    border = Border(left=border_style, right=border_style, top=border_style, bottom=border_style)
    alignment = Alignment(horizontal='center', vertical='center')

    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for cell in ws[1]:  # Header row
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
        
        # Adjust column widths and apply borders
        for col in ws.columns:
            max_length = 0
            column_letter = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
                cell.border = border
                if cell.row != 1:
                    cell.alignment = alignment
            
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column_letter].width = adjusted_width
            
    wb.save(excel_file)
    print(f"Excel file created at {excel_file} with formatting.")

def json_to_csv(json_file, csv_file):
    """Convert flat JSON data to CSV. Will extract batting scorecard as an example."""
    data = read_json(json_file)
    if not data:
        return
    
    ind_bat_df = pd.DataFrame(data['batting_stats']['India'])
    ind_bat_df['Team'] = 'India'
    aus_bat_df = pd.DataFrame(data['batting_stats']['Australia'])
    aus_bat_df['Team'] = 'Australia'
    
    combined_batting = pd.concat([ind_bat_df, aus_bat_df], ignore_index=True)
    combined_batting.to_csv(csv_file, index=False)
    print(f"Saved Batting Stats CSV to {csv_file}")

def csv_to_json(csv_file, json_file):
    """Read a CSV file and convert it into JSON."""
    if not os.path.exists(csv_file):
        print(f"Error: {csv_file} does not exist.")
        return
    df = pd.read_csv(csv_file)
    # Convert dataframe to a dictionary records list
    data = df.to_dict(orient='records')
    # Save as JSON
    save_json({"imported_data": data}, json_file)
    print(f"Converted {csv_file} to JSON and saved at {json_file}")
